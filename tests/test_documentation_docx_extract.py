"""
================================================================================
Module:     tests/test_documentation_docx_extract.py
Project:    Electronics Diagnostic Tool Hub / ESR Tester (Windows)
Versie:     1.0.0
Datum:      2026-10-09
Auteur:     Bart Bossuyt

Doel:       Regressietests voor docx_extract: metadata, tekstextractie,
            hashing, kopiëren naar sources/ en de .oud-rollback-helpers.

Wijzigingen:
  v1.0.0 (2026-10-09)  Eerste versie (fase 6A).
================================================================================
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.documentation.docx_extract import (
    DocxExtractError,
    DocxMetadata,
    compute_file_hash,
    copy_docx_to_sources,
    extract_docx_metadata,
    extract_docx_text,
    rename_existing_docx_to_old,
    restore_old_docx,
)


def _maak_docx(pad: Path, *, titel: str = "Testdocument", auteur: str = "Tester") -> Path:
    document = Document()
    document.add_paragraph("Eerste paragraaf.")
    document.add_paragraph("Tweede paragraaf.")
    document.core_properties.title = titel
    document.core_properties.author = auteur
    document.save(str(pad))
    return pad


# ---------------------------------------------------------------- metadata

def test_extract_metadata_gelukkig(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx", titel="Mijn titel", auteur="Auteur")
    meta = extract_docx_metadata(pad)

    assert isinstance(meta, DocxMetadata)
    assert meta.title == "Mijn titel"
    assert meta.author == "Auteur"
    assert meta.paragraph_count == 2


def test_extract_metadata_ontbrekend_bestand(tmp_path):
    with pytest.raises(DocxExtractError):
        extract_docx_metadata(tmp_path / "bestaat-niet.docx")


def test_extract_metadata_geen_bestand(tmp_path):
    map_pad = tmp_path / "map"
    map_pad.mkdir()
    with pytest.raises(DocxExtractError):
        extract_docx_metadata(map_pad)


def test_extract_metadata_corrupt_bestand(tmp_path):
    ongeldig = tmp_path / "kapot.docx"
    ongeldig.write_text("geen docx", encoding="utf-8")
    with pytest.raises(DocxExtractError):
        extract_docx_metadata(ongeldig)


def test_extract_metadata_zonder_core_properties(tmp_path):
    """Een minimale .docx zonder metadata levert lege velden."""
    document = Document()
    document.add_paragraph("Inhoud")
    pad = tmp_path / "leeg.docx"
    document.save(str(pad))

    meta = extract_docx_metadata(pad)
    assert meta.paragraph_count == 1
    # core_properties zijn optioneel; geen crash bij lege waarden.


# ---------------------------------------------------------------- tekst

def test_extract_text_gelukkig(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    tekst = extract_docx_text(pad)
    assert "Eerste paragraaf." in tekst
    assert "Tweede paragraaf." in tekst


def test_extract_text_ontbrekend_bestand(tmp_path):
    with pytest.raises(DocxExtractError):
        extract_docx_text(tmp_path / "bestaat-niet.docx")


# ---------------------------------------------------------------- hash

def test_compute_hash_deterministisch(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    h1 = compute_file_hash(pad)
    h2 = compute_file_hash(pad)
    assert h1 == h2
    assert len(h1) == 64  # SHA-256 hex


def test_compute_hash_ontbrekend_bestand(tmp_path):
    with pytest.raises(DocxExtractError):
        compute_file_hash(tmp_path / "bestaat-niet.docx")


# ---------------------------------------------------------------- kopiëren

def test_copy_naar_sources(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    sources = tmp_path / "sources"

    doel = copy_docx_to_sources(pad, "source-1", sources_dir=sources)
    assert doel.exists()
    assert doel.name == "source-1.docx"


def test_copy_weigert_bestaand_doel(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "source-1.docx").write_text("bezet", encoding="utf-8")

    with pytest.raises(DocxExtractError):
        copy_docx_to_sources(pad, "source-1", sources_dir=sources)


def test_copy_lege_source_id(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    with pytest.raises(DocxExtractError):
        copy_docx_to_sources(pad, "", sources_dir=tmp_path / "sources")


# ---------------------------------------------------------------- rollback

def test_rename_naar_oud(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    oud = rename_existing_docx_to_old(pad, timestamp="20261009")
    assert oud is not None
    assert oud.exists()
    assert not pad.exists()


def test_rename_naar_oud_ontbrekend_bestand(tmp_path):
    result = rename_existing_docx_to_old(
        tmp_path / "bestaat-niet.docx", timestamp="20261009"
    )
    assert result is None


def test_restore_old(tmp_path):
    pad = _maak_docx(tmp_path / "doc.docx")
    oud = rename_existing_docx_to_old(pad, timestamp="20261009")
    assert oud is not None

    restore_old_docx(oud, pad)
    assert pad.exists()
    assert not oud.exists()