"""
================================================================================
Module:     tests/test_import_wizard_dialog.py
Project:    Electronics Diagnostic Tool Hub / ESR Tester (Windows)
Versie:     1.7.0
Datum:      2026-10-09
Auteur:     Bart Bossuyt

Doel:       GUI-regressietests voor de import-wizard. Netwerk, PDF-parsing
            en ImportService worden gemockt of in tmp_path geïsoleerd.
            Geen echte HTTP-verzoeken. Sinds v1.1.0 duplicate-popup-flow.
            Sinds v1.2.0 startmap-logica. Sinds v1.2.1 titelveld-gedrag.
            Sinds v1.3.0 metadata-sectie (5D'.2b).
            Sinds v1.4.0 UX-verfijning: QDateEdit + "Datum onbekend" +
            Ctrl+D + uppercase voor identificatievelden (5D'.2d).
            Sinds v1.5.0 gebruikt de wizard de gedeelde helpers uit
            _metadata_form_helpers; tests importeren daar ook
            naar_uppercase uit.
            Sinds v1.6.0 ondersteunt de wizard ook Word (.docx) en
            Excel (.xlsx); nieuwe sectie v1.6.0 test de nieuwe types,
            paneelwissels, samenvatting en import.
            Sinds v1.7.0 toont de wizard een succesmelding na een
            geslaagde import. De test patcht _toon_succesmelding op de
            dialoog-instantie (QMessageBox.information is een
            C++-staticmethod die zich niet betrouwbaar laat patchen en
            de test zou laten hangen op een modale dialoog).

Wijzigingen:
  v1.0.0 (2026-10-06)  Eerste versie.
  v1.0.1 (2026-10-06)  isVisible() vervangen door isHidden().
  v1.0.2 (2026-10-07)  Titel-voorstel bij PDF is nu bestandsnaam.
  v1.1.0 (2026-10-07)  Tests voor duplicate-popup.
  v1.2.0 (2026-10-07)  Tests voor startmap en laatste_importmap.
  v1.2.1 (2026-10-07)  Tests voor titelveld-gedrag bij meerdere PDF-keuzes.
  v1.3.0 (2026-10-07)  Tests voor metadata-sectie.
  v1.4.0 (2026-10-07)  Tests voor QDateEdit, "Datum onbekend", Ctrl+D
                       en uppercase-conversie.
  v1.5.0 (2026-10-07)  Fix: naar_uppercase wordt geïmporteerd uit
                       _metadata_form_helpers (verplaatst in wizard
                       v1.5.0). Alle aanroepen hernoemd.
  v1.6.0 (2026-10-09)  Fase 6B: tests voor Word- en Excel-import in
                       de wizard. Nieuwe tests gebruiken findData in
                       plaats van hardcoded categorie-indexen.
  v1.7.0 (2026-10-09)  Fase 6C: test voor succesmelding na import.
                       Patch _toon_succesmelding op de dialoog-instantie
                       in plaats van QMessageBox.information; dat laatste
                       is een C++-staticmethod die zich niet betrouwbaar
                       laat monkeypatchen en de test zou laten hangen.
================================================================================
"""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import QApplication

import app.gui.dialogs.import_wizard_dialog as wizard_module
from app.config.settings import AppInstellingen
from app.documentation.import_models import (
    DuplicateAction,
    DuplicateMatch,
    ImportResult,
    ImportSource,
    ImportSourceType,
    ImportStatus,
)
from app.documentation.import_service import ImportService
from app.gui.dialogs._metadata_form_helpers import naar_uppercase
from app.gui.dialogs.import_wizard_dialog import ImportWizardDialog


def _app():
    return QApplication.instance() or QApplication([])


def _maak_eenvoudige_pdf(pad: Path, *, titel: str = "Testdocument") -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": titel})
    with pad.open("wb") as handle:
        writer.write(handle)
    return pad


def _maak_docx(pad: Path, *, titel: str = "Testdocument") -> Path:
    from docx import Document

    document = Document()
    document.add_paragraph("Inhoud")
    document.core_properties.title = titel
    document.save(str(pad))
    return pad


def _maak_xlsx(pad: Path, *, titel: str = "Testwerkmap") -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Inhoud"
    wb.properties.title = titel
    wb.save(str(pad))
    return pad


# ============================================================================
# Basis
# ============================================================================

def test_dialoog_opent_met_pdf_geselecteerd(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.radio_pdf.isChecked() is True
    assert dialog.radio_word.isChecked() is False
    assert dialog.radio_excel.isChecked() is False
    assert dialog.radio_url.isChecked() is False
    assert dialog.pdf_panel.isHidden() is False
    assert dialog.word_panel.isHidden() is True
    assert dialog.excel_panel.isHidden() is True
    assert dialog.url_panel.isHidden() is True
    assert dialog.import_btn.isEnabled() is False


def test_type_wissel_toont_url_paneel(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.radio_url.setChecked(True)

    assert dialog.url_panel.isHidden() is False
    assert dialog.pdf_panel.isHidden() is True
    assert dialog.word_panel.isHidden() is True
    assert dialog.excel_panel.isHidden() is True


# ============================================================================
# PDF kiezen
# ============================================================================

def test_kies_pdf_vult_samenvatting(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    dialog._pick_pdf()

    assert dialog._pdf_pad == pdf_pad
    assert dialog._pdf_meta_titel == "Mijn PDF"
    assert dialog.title_edit.text() == "doc"
    assert dialog.import_btn.isEnabled() is True
    assert dialog._pdf_hash is not None


def test_pdf_titel_komt_uit_bestandsnaam_niet_uit_metadata(
    tmp_path, monkeypatch
):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(
        tmp_path / "CHONGCD11XSERIES.pdf",
        titel="Microsoft Word - CD11X n3-4 105.doc",
    )

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    dialog._pick_pdf()

    assert dialog.title_edit.text() == "CHONGCD11XSERIES"
    assert dialog._pdf_meta_titel == "Microsoft Word - CD11X n3-4 105.doc"


def test_kies_ongeldige_pdf_toont_fout(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    ongeldig = tmp_path / "geen.pdf"
    ongeldig.write_text("geen pdf", encoding="utf-8")

    from PySide6.QtWidgets import QFileDialog, QMessageBox
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(ongeldig), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    gewaarschuwd = {"count": 0}
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        staticmethod(
            lambda *a, **kw: gewaarschuwd.__setitem__(
                "count", gewaarschuwd["count"] + 1
            )
        ),
    )

    dialog._pick_pdf()

    assert gewaarschuwd["count"] == 1
    assert dialog._pdf_pad is None
    assert dialog.import_btn.isEnabled() is False


# ============================================================================
# URL
# ============================================================================

def test_url_metadata_ophalen_vult_titel(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_url.setChecked(True)

    from app.documentation.url_fetch import UrlMetadata
    nep_meta = UrlMetadata(
        url="https://example.com",
        fetched_at=1,
        http_status=200,
        title="Paginatitel",
        og_title=None,
    )
    monkeypatch.setattr(
        wizard_module, "fetch_url_metadata", lambda url: nep_meta
    )

    dialog.url_edit.setText("https://example.com")
    dialog._fetch_url_metadata()

    assert dialog._url == "https://example.com"
    assert dialog._url_meta is nep_meta
    assert dialog.title_edit.text() == "Paginatitel"
    assert dialog.import_btn.isEnabled() is True


def test_url_metadata_fout_toont_waarschuwing(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_url.setChecked(True)
    dialog.url_edit.setText("https://example.com")

    from app.documentation.url_fetch import UrlFetchError

    def _faal(url):
        raise UrlFetchError("timeout")

    monkeypatch.setattr(wizard_module, "fetch_url_metadata", _faal)

    from PySide6.QtWidgets import QMessageBox
    gewaarschuwd = {"count": 0}
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        staticmethod(
            lambda *a, **kw: gewaarschuwd.__setitem__(
                "count", gewaarschuwd["count"] + 1
            )
        ),
    )

    dialog._fetch_url_metadata()

    assert gewaarschuwd["count"] == 1
    assert dialog._url_meta is None
    assert dialog.import_btn.isEnabled() is False


# ============================================================================
# Importeren
# ============================================================================

def test_importeren_pdf_emit_signal(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    # Voorkom dat de succesmelding de test laat hangen.
    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    ontvangen = {"source_id": None}
    dialog.import_completed.connect(
        lambda sid: ontvangen.__setitem__("source_id", sid)
    )

    dialog._perform_import()

    assert ontvangen["source_id"] is not None
    assert len(service.list_sources()) == 1


def test_importeren_fout_toont_waarschuwing(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    from app.documentation.pdf_extract import PdfExtractError

    def _faal(*a, **kw):
        raise PdfExtractError("test-fout")

    monkeypatch.setattr(wizard_module, "import_pdf", _faal)

    from PySide6.QtWidgets import QMessageBox
    gewaarschuwd = {"count": 0}
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        staticmethod(
            lambda *a, **kw: gewaarschuwd.__setitem__(
                "count", gewaarschuwd["count"] + 1
            )
        ),
    )

    dialog._perform_import()

    assert gewaarschuwd["count"] == 1
    assert len(service.list_sources()) == 0


# ============================================================================
# Taalwissel
# ============================================================================

def test_titels_komen_uit_i18n(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog_nl = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog_en = ImportWizardDialog(taal="en_US", import_service=service)

    assert dialog_nl.radio_pdf.text() == "PDF-bestand"
    assert dialog_en.radio_pdf.text() == "PDF file"
    assert dialog_nl.cancel_btn.text() == "Annuleren"
    assert dialog_en.cancel_btn.text() == "Cancel"


# ============================================================================
# v1.1.0 — duplicate-popup-flow
# ============================================================================

def _maak_match(
    *,
    source_type: ImportSourceType = ImportSourceType.PDF,
    match_type: str = "file_hash",
) -> DuplicateMatch:
    if source_type is ImportSourceType.PDF:
        bestaande = ImportSource(
            source_id="abc-123",
            source_type=ImportSourceType.PDF,
            title="Bestaand document",
            imported_at=1_700_000_000_000,
            imported_by="tester",
            status=ImportStatus.CONCEPT,
            original_filename="oud.pdf",
            file_hash="deadbeef",
        )
    else:
        bestaande = ImportSource(
            source_id="abc-456",
            source_type=ImportSourceType.URL,
            title="Bestaande URL",
            imported_at=1_700_000_000_000,
            imported_by="tester",
            status=ImportStatus.CONCEPT,
            source_url="https://example.com",
        )
    return DuplicateMatch(bestaande=bestaande, match_type=match_type)


def test_duplicate_check_geen_match_voor_pdf_zonder_keuze(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    assert dialog._check_duplicate() is None


def test_duplicate_check_vindt_bestaande_pdf(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    from app.documentation.pdf_import import import_pdf
    import_pdf(
        pdf_pad,
        import_service=service,
        sources_dir=tmp_path / "sources",
    )

    match = dialog._check_duplicate()
    assert match is not None
    assert match.match_type == "file_hash"


def test_wizard_keep_bij_match_doet_niets(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    from app.documentation.pdf_import import import_pdf
    eerste = import_pdf(
        pdf_pad,
        import_service=service,
        sources_dir=tmp_path / "sources",
    )
    aantal_voor = len(service.list_sources())

    from app.gui.dialogs import duplicate_source_dialog as dup_module
    monkeypatch.setattr(
        dup_module.DuplicateSourceDialog,
        "vraag_actie",
        staticmethod(lambda **kw: DuplicateAction.KEEP),
    )

    dialog._perform_import()

    assert len(service.list_sources()) == aantal_voor
    assert service.get(eerste.source.source_id).status is ImportStatus.CONCEPT


def test_wizard_annuleren_in_popup_doet_niets(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    from app.documentation.pdf_import import import_pdf
    import_pdf(
        pdf_pad,
        import_service=service,
        sources_dir=tmp_path / "sources",
    )
    aantal_voor = len(service.list_sources())

    from app.gui.dialogs import duplicate_source_dialog as dup_module
    monkeypatch.setattr(
        dup_module.DuplicateSourceDialog,
        "vraag_actie",
        staticmethod(lambda **kw: None),
    )

    dialog._perform_import()

    assert len(service.list_sources()) == aantal_voor


def test_wizard_overwrite_bij_match(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    from app.documentation.pdf_import import import_pdf
    eerste = import_pdf(
        pdf_pad,
        import_service=service,
        sources_dir=tmp_path / "sources",
    )

    from app.gui.dialogs import duplicate_source_dialog as dup_module
    monkeypatch.setattr(
        dup_module.DuplicateSourceDialog,
        "vraag_actie",
        staticmethod(lambda **kw: DuplicateAction.OVERWRITE),
    )

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    ontvangen = {"source_id": None}
    dialog.import_completed.connect(
        lambda sid: ontvangen.__setitem__("source_id", sid)
    )

    dialog._perform_import()

    assert ontvangen["source_id"] == eerste.source.source_id
    assert len(service.list_sources()) == 1


# ============================================================================
# v1.2.0 — startmap-logica (fase 5D'.3)
# ============================================================================

def test_start_map_gebruikt_laatste_importmap_indien_geldig(
    tmp_path, monkeypatch
):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    laatste_map = tmp_path / "laatste"
    laatste_map.mkdir()

    basis = AppInstellingen()
    from dataclasses import replace as dc_replace

    instellingen = dc_replace(
        basis,
        algemeen=dc_replace(
            basis.algemeen,
            laatste_importmap=str(laatste_map),
            standaard_importmap=str(tmp_path / "standaard"),
        ),
    )

    start = dialog._start_map_voor_pdf(instellingen)
    assert start == str(laatste_map)


def test_start_map_valt_terug_op_standaardmap_bij_ongeldige_laatste(
    tmp_path, monkeypatch
):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    standaard_map = tmp_path / "standaard"
    standaard_map.mkdir()

    basis = AppInstellingen()
    from dataclasses import replace as dc_replace

    instellingen = dc_replace(
        basis,
        algemeen=dc_replace(
            basis.algemeen,
            laatste_importmap=str(tmp_path / "bestaat-niet"),
            standaard_importmap=str(standaard_map),
        ),
    )

    start = dialog._start_map_voor_pdf(instellingen)
    assert start == str(standaard_map)


def test_start_map_valt_terug_op_leeg_bij_geen_geldige_mappen(
    tmp_path
):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    basis = AppInstellingen()
    from dataclasses import replace as dc_replace

    instellingen = dc_replace(
        basis,
        algemeen=dc_replace(
            basis.algemeen,
            laatste_importmap=str(tmp_path / "bestaat-niet-1"),
            standaard_importmap=str(tmp_path / "bestaat-niet-2"),
        ),
    )

    start = dialog._start_map_voor_pdf(instellingen)
    assert start == ""


def test_pick_pdf_start_in_laatste_importmap(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    laatste_map = tmp_path / "laatste"
    laatste_map.mkdir()
    pdf_pad = _maak_eenvoudige_pdf(laatste_map / "doc.pdf")

    basis = AppInstellingen()
    from dataclasses import replace as dc_replace

    instellingen = dc_replace(
        basis,
        algemeen=dc_replace(
            basis.algemeen,
            laatste_importmap=str(laatste_map),
        ),
    )
    monkeypatch.setattr(
        wizard_module, "laad_instellingen", lambda: instellingen
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    ontvangen_start = {"pad": None}

    from PySide6.QtWidgets import QFileDialog

    def _fake_get_open(parent, titel, start, filter_):
        ontvangen_start["pad"] = start
        return (str(pdf_pad), "PDF (*.pdf)")

    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", staticmethod(_fake_get_open)
    )

    dialog._pick_pdf()

    assert ontvangen_start["pad"] == str(laatste_map)


def test_pick_pdf_onthoudt_nieuwe_map(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    andere_map = tmp_path / "andere"
    andere_map.mkdir()
    pdf_pad = _maak_eenvoudige_pdf(andere_map / "doc.pdf")

    basis = AppInstellingen()
    monkeypatch.setattr(
        wizard_module, "laad_instellingen", lambda: basis
    )

    opgeslagen = {"instellingen": None}

    def _fake_opslaan(inst):
        opgeslagen["instellingen"] = inst

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", _fake_opslaan)

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )

    dialog._pick_pdf()

    assert opgeslagen["instellingen"] is not None
    assert (
        opgeslagen["instellingen"].algemeen.laatste_importmap
        == str(andere_map)
    )


def test_pick_pdf_annuleren_wijzigt_niets(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    basis = AppInstellingen()
    monkeypatch.setattr(
        wizard_module, "laad_instellingen", lambda: basis
    )

    opgeslagen = {"count": 0}

    def _fake_opslaan(inst):
        opgeslagen["count"] += 1

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", _fake_opslaan)

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: ("", "")),
    )

    dialog._pick_pdf()

    assert opgeslagen["count"] == 0
    assert dialog._pdf_pad is None


def test_pick_pdf_zelfde_map_slaat_niet_opnieuw_op(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")

    basis = AppInstellingen()
    from dataclasses import replace as dc_replace

    instellingen = dc_replace(
        basis,
        algemeen=dc_replace(
            basis.algemeen,
            laatste_importmap=str(tmp_path),
        ),
    )
    monkeypatch.setattr(
        wizard_module, "laad_instellingen", lambda: instellingen
    )

    opgeslagen = {"count": 0}

    def _fake_opslaan(inst):
        opgeslagen["count"] += 1

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", _fake_opslaan)

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )

    dialog._pick_pdf()

    assert opgeslagen["count"] == 0


# ============================================================================
# v1.2.1 — titelveld-gedrag bij meerdere PDF-keuzes
# ============================================================================

def test_titel_wisselt_bij_nieuwe_pdf_als_automatisch(tmp_path, monkeypatch):
    """Twee PDF-keuzes na elkaar: titel volgt de tweede keuze."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_a = _maak_eenvoudige_pdf(tmp_path / "eerste.pdf")
    pdf_b = _maak_eenvoudige_pdf(tmp_path / "tweede.pdf")

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    from PySide6.QtWidgets import QFileDialog

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_a), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    assert dialog.title_edit.text() == "eerste"
    assert dialog._titel_automatisch is True

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_b), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    assert dialog.title_edit.text() == "tweede"


def test_titel_blijft_staan_bij_nieuwe_pdf_na_handmatige_bewerking(
    tmp_path, monkeypatch
):
    """Als de gebruiker de titel handmatig typt, blijft die staan."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_a = _maak_eenvoudige_pdf(tmp_path / "eerste.pdf")
    pdf_b = _maak_eenvoudige_pdf(tmp_path / "tweede.pdf")

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    from PySide6.QtWidgets import QFileDialog

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_a), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    assert dialog.title_edit.text() == "eerste"

    dialog._on_titel_handmatig_bewerkt("Mijn eigen titel")
    dialog.title_edit.setText("Mijn eigen titel")
    assert dialog._titel_automatisch is False

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_b), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    assert dialog.title_edit.text() == "Mijn eigen titel"


def test_samenvatting_wisselt_bij_nieuwe_pdf(tmp_path, monkeypatch):
    """Samenvatting toont altijd de laatste PDF, ongeacht titel-state."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_a = _maak_eenvoudige_pdf(tmp_path / "eerste.pdf")
    pdf_b = _maak_eenvoudige_pdf(tmp_path / "tweede.pdf")

    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    from PySide6.QtWidgets import QFileDialog

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_a), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    html_a = dialog.summary_browser.toHtml()
    assert "eerste.pdf" in html_a

    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_b), "PDF (*.pdf)")),
    )
    dialog._pick_pdf()
    html_b = dialog.summary_browser.toHtml()
    assert "tweede.pdf" in html_b
    assert "eerste.pdf" not in html_b


# ============================================================================
# v1.3.0 — metadata-sectie (fase 5D'.2b)
# ============================================================================

def test_metadata_sectie_aanwezig_bij_opstart(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.metadata_label.text() == "Metadata"
    assert dialog.categorie_label.text() == "Categorie"
    assert dialog.fabrikant_label.text() == "Fabrikant"
    assert dialog.serie_label.text() == "Serie"
    assert dialog.partnummer_label.text() == "Partnummer"
    assert dialog.documentversie_label.text() == "Documentversie"
    assert dialog.documentdatum_label.text() == "Documentdatum"
    assert dialog.notities_label.text() == "Notities"


def test_categorie_dropdown_gevuld_met_zeven_categorieen(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.categorie_combo.count() == 7
    assert dialog.categorie_combo.currentData() == "DATASHEET"
    assert dialog.categorie_combo.currentText() == "Fabrikantdatasheet"


def test_categorie_dropdown_en_labels(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="en_US", import_service=service)

    assert dialog.categorie_combo.count() == 7
    assert dialog.categorie_combo.currentData() == "DATASHEET"
    assert dialog.categorie_combo.currentText() == "Manufacturer datasheet"


def test_huidige_categorie_leest_itemdata(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    index = dialog.categorie_combo.findData("MANUAL")
    dialog.categorie_combo.setCurrentIndex(index)
    assert dialog._huidige_categorie() == "MANUAL"


def test_metadata_uit_formulier_bevat_geen_notes(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.fabrikant_edit.setText("Panasonic")
    dialog.serie_edit.setText("FR")
    dialog.partnummer_edit.setText("FR-123")
    dialog.documentversie_edit.setText("1.2")

    metadata = dialog._metadata_uit_formulier()
    assert "notes" not in metadata
    assert metadata["category"] == "DATASHEET"
    assert metadata["manufacturer"] == "Panasonic"
    assert metadata["series"] == "FR"
    assert metadata["part_number"] == "FR-123"
    assert metadata["document_version"] == "1.2"


def test_notities_uit_formulier_leeg_geeft_none(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog._notities_uit_formulier() is None
    dialog.notities_edit.setText("   ")
    assert dialog._notities_uit_formulier() is None
    dialog.notities_edit.setText("noot")
    assert dialog._notities_uit_formulier() == "noot"


def test_importeren_met_metadata(tmp_path, monkeypatch):
    """Metadata uit de wizard komt in de gebruikerscatalogus."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    index = dialog.categorie_combo.findData("MANUAL")
    dialog.categorie_combo.setCurrentIndex(index)
    dialog.fabrikant_edit.setText("Panasonic")
    dialog.serie_edit.setText("FR")
    dialog.partnummer_edit.setText("FR-123")
    dialog.documentversie_edit.setText("1.2")
    dialog.notities_edit.setText("noot")

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    s = items[0]
    assert s.category == "MANUAL"
    assert s.manufacturer == "Panasonic"
    assert s.series == "FR"
    assert s.part_number == "FR-123"
    assert s.document_version == "1.2"
    assert s.notes == "noot"


def test_importeren_zonder_metadata_laat_velden_none(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    dialog.datum_onbekend_checkbox.setChecked(True)

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    s = items[0]
    assert s.category == "DATASHEET"
    assert s.manufacturer is None
    assert s.series is None
    assert s.part_number is None
    assert s.document_version is None
    assert s.document_date is None
    assert s.notes is None


def test_taalwissel_behoudt_categorie_selectie(tmp_path, monkeypatch):
    """Bij een taalwissel blijft de gekozen categorie staan."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    index = dialog.categorie_combo.findData("MANUAL")
    dialog.categorie_combo.setCurrentIndex(index)
    dialog.taal = "en_US"
    dialog._apply_language()

    assert dialog.categorie_combo.currentData() == "MANUAL"
    assert dialog.categorie_combo.currentText() == "Manual"


# ============================================================================
# v1.4.0 — UX-verfijning (fase 5D'.2d)
# ============================================================================

def test_documentdatum_is_qdateedit(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    from PySide6.QtWidgets import QDateEdit
    assert isinstance(dialog.documentdatum_edit, QDateEdit)
    assert dialog.documentdatum_edit.calendarPopup() is True
    assert dialog.documentdatum_edit.displayFormat() == "yyyy-MM-dd"


def test_documentdatum_default_vandaag(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.documentdatum_edit.date() == QDate.currentDate()


def test_documentdatum_onbekend_checkbox_default_uit(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.datum_onbekend_checkbox.isChecked() is False
    assert dialog.documentdatum_edit.isEnabled() is True


def test_documentdatum_onbekend_disablet_veld(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.datum_onbekend_checkbox.setChecked(True)
    assert dialog.documentdatum_edit.isEnabled() is False

    dialog.datum_onbekend_checkbox.setChecked(False)
    assert dialog.documentdatum_edit.isEnabled() is True


def test_documentdatum_waarde_iso_formaat(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.documentdatum_edit.setDate(QDate(2024, 1, 31))
    assert dialog._documentdatum_waarde() == "2024-01-31"


def test_documentdatum_onbekend_geeft_none(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.datum_onbekend_checkbox.setChecked(True)
    assert dialog._documentdatum_waarde() is None


def test_documentdatum_ctrl_d_zet_vandaag(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.documentdatum_edit.setDate(QDate(2000, 1, 1))
    dialog._zet_datum_op_vandaag()
    assert dialog.documentdatum_edit.date() == QDate.currentDate()


def test_documentdatum_ctrl_d_negeert_onbekend(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.documentdatum_edit.setDate(QDate(2000, 1, 1))
    dialog.datum_onbekend_checkbox.setChecked(True)
    dialog._zet_datum_op_vandaag()
    assert dialog.documentdatum_edit.date() == QDate(2000, 1, 1)


def test_uppercase_fabrikant(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.fabrikant_edit.setText("panasonic")
    naar_uppercase(dialog.fabrikant_edit)
    assert dialog.fabrikant_edit.text() == "PANASONIC"


def test_uppercase_serie(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.serie_edit.setText("fr-serie")
    naar_uppercase(dialog.serie_edit)
    assert dialog.serie_edit.text() == "FR-SERIE"


def test_uppercase_partnummer(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.partnummer_edit.setText("fr-123")
    naar_uppercase(dialog.partnummer_edit)
    assert dialog.partnummer_edit.text() == "FR-123"


def test_uppercase_documentversie(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.documentversie_edit.setText("v1.2")
    naar_uppercase(dialog.documentversie_edit)
    assert dialog.documentversie_edit.text() == "V1.2"


def test_titel_blijft_gemengd(tmp_path):
    """Titel wordt niet naar uppercase geconverteerd."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.title_edit.setText("Measure ESR with an ESR meter")
    assert dialog.title_edit.text() == "Measure ESR with an ESR meter"


def test_notities_blijft_gemengd(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.notities_edit.setText("Test notitie met Gemengde Case")
    assert dialog.notities_edit.text() == "Test notitie met Gemengde Case"


def test_uppercase_behoudt_cursorpositie(tmp_path):
    """Cursorpositie blijft behouden na uppercase-conversie."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.fabrikant_edit.setText("abc")
    dialog.fabrikant_edit.setCursorPosition(2)
    naar_uppercase(dialog.fabrikant_edit)
    assert dialog.fabrikant_edit.text() == "ABC"
    assert dialog.fabrikant_edit.cursorPosition() == 2


def test_uppercase_lege_tekst_blijft_leeg(tmp_path):
    """Een leeg veld blijft leeg; geen crash."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.fabrikant_edit.setText("")
    naar_uppercase(dialog.fabrikant_edit)
    assert dialog.fabrikant_edit.text() == ""


def test_uppercase_al_uppercase_doet_niets(tmp_path):
    """Als de tekst al uppercase is, gebeurt er niets."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.fabrikant_edit.setText("PANASONIC")
    dialog.fabrikant_edit.setCursorPosition(3)
    naar_uppercase(dialog.fabrikant_edit)
    assert dialog.fabrikant_edit.text() == "PANASONIC"
    assert dialog.fabrikant_edit.cursorPosition() == 3


def test_importeren_met_datum_onbekend(tmp_path, monkeypatch):
    """Met 'Datum onbekend' aangevinkt: document_date is None."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    dialog.datum_onbekend_checkbox.setChecked(True)
    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)
    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    assert items[0].document_date is None


def test_importeren_met_datum_ingevuld(tmp_path, monkeypatch):
    """Met een datum: document_date is ISO-string."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    dialog.documentdatum_edit.setDate(QDate(2024, 1, 31))
    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)
    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    assert items[0].document_date == "2024-01-31"


# ============================================================================
# v1.6.0 — Word en Excel (fase 6B)
# ============================================================================

def test_type_wissel_toont_word_paneel(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.radio_word.setChecked(True)

    assert dialog.word_panel.isHidden() is False
    assert dialog.pdf_panel.isHidden() is True
    assert dialog.excel_panel.isHidden() is True
    assert dialog.url_panel.isHidden() is True


def test_type_wissel_toont_excel_paneel(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    dialog.radio_excel.setChecked(True)

    assert dialog.excel_panel.isHidden() is False
    assert dialog.pdf_panel.isHidden() is True
    assert dialog.word_panel.isHidden() is True
    assert dialog.url_panel.isHidden() is True


def test_kies_docx_vult_samenvatting(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_word.setChecked(True)

    docx_pad = _maak_docx(tmp_path / "doc.docx", titel="Mijn Word")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(docx_pad), "Word (*.docx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    dialog._pick_docx()

    assert dialog._docx_pad == docx_pad
    assert dialog._docx_meta_titel == "Mijn Word"
    assert dialog.title_edit.text() == "doc"
    assert dialog.import_btn.isEnabled() is True
    assert dialog._docx_hash is not None


def test_kies_xlsx_vult_samenvatting(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_excel.setChecked(True)

    xlsx_pad = _maak_xlsx(tmp_path / "doc.xlsx", titel="Mijn Werkmap")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(xlsx_pad), "Excel (*.xlsx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)

    dialog._pick_xlsx()

    assert dialog._xlsx_pad == xlsx_pad
    assert dialog._xlsx_meta_titel == "Mijn Werkmap"
    assert dialog.title_edit.text() == "doc"
    assert dialog.import_btn.isEnabled() is True
    assert dialog._xlsx_hash is not None


def test_kies_ongeldige_docx_toont_fout(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_word.setChecked(True)

    ongeldig = tmp_path / "kapot.docx"
    ongeldig.write_text("geen docx", encoding="utf-8")

    from PySide6.QtWidgets import QFileDialog, QMessageBox
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(ongeldig), "Word (*.docx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    gewaarschuwd = {"count": 0}
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        staticmethod(
            lambda *a, **kw: gewaarschuwd.__setitem__(
                "count", gewaarschuwd["count"] + 1
            )
        ),
    )

    dialog._pick_docx()

    assert gewaarschuwd["count"] == 1
    assert dialog._docx_pad is None
    assert dialog.import_btn.isEnabled() is False


def test_kies_ongeldige_xlsx_toont_fout(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_excel.setChecked(True)

    ongeldig = tmp_path / "kapot.xlsx"
    ongeldig.write_text("geen xlsx", encoding="utf-8")

    from PySide6.QtWidgets import QFileDialog, QMessageBox
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(ongeldig), "Excel (*.xlsx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    gewaarschuwd = {"count": 0}
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        staticmethod(
            lambda *a, **kw: gewaarschuwd.__setitem__(
                "count", gewaarschuwd["count"] + 1
            )
        ),
    )

    dialog._pick_xlsx()

    assert gewaarschuwd["count"] == 1
    assert dialog._xlsx_pad is None
    assert dialog.import_btn.isEnabled() is False


def test_importeren_word_emit_signal(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_word.setChecked(True)

    docx_pad = _maak_docx(tmp_path / "doc.docx")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(docx_pad), "Word (*.docx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_docx()

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    ontvangen = {"source_id": None}
    dialog.import_completed.connect(
        lambda sid: ontvangen.__setitem__("source_id", sid)
    )

    dialog._perform_import()

    assert ontvangen["source_id"] is not None
    items = service.list_sources()
    assert len(items) == 1
    assert items[0].source_type.value == "docx"


def test_importeren_excel_emit_signal(tmp_path, monkeypatch):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_excel.setChecked(True)

    xlsx_pad = _maak_xlsx(tmp_path / "doc.xlsx")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(xlsx_pad), "Excel (*.xlsx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_xlsx()

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    ontvangen = {"source_id": None}
    dialog.import_completed.connect(
        lambda sid: ontvangen.__setitem__("source_id", sid)
    )

    dialog._perform_import()

    assert ontvangen["source_id"] is not None
    items = service.list_sources()
    assert len(items) == 1
    assert items[0].source_type.value == "xlsx"


def test_importeren_word_met_metadata(tmp_path, monkeypatch):
    """Metadata uit de wizard komt in de gebruikerscatalogus voor Word."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_word.setChecked(True)

    docx_pad = _maak_docx(tmp_path / "doc.docx")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(docx_pad), "Word (*.docx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_docx()

    index = dialog.categorie_combo.findData("MANUAL")
    dialog.categorie_combo.setCurrentIndex(index)
    dialog.fabrikant_edit.setText("CHONG")
    dialog.notities_edit.setText("Word noot")

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    s = items[0]
    assert s.category == "MANUAL"
    assert s.manufacturer == "CHONG"
    assert s.notes == "Word noot"
    assert s.source_type.value == "docx"


def test_importeren_excel_met_metadata(tmp_path, monkeypatch):
    """Metadata uit de wizard komt in de gebruikerscatalogus voor Excel."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_excel.setChecked(True)

    xlsx_pad = _maak_xlsx(tmp_path / "doc.xlsx")

    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(xlsx_pad), "Excel (*.xlsx)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_xlsx()

    index = dialog.categorie_combo.findData("REFERENCE_TABLE")
    dialog.categorie_combo.setCurrentIndex(index)
    dialog.fabrikant_edit.setText("TDK")

    monkeypatch.setattr(dialog, "_toon_succesmelding", lambda titel: None)

    dialog._perform_import()

    items = service.list_sources()
    assert len(items) == 1
    s = items[0]
    assert s.category == "REFERENCE_TABLE"
    assert s.manufacturer == "TDK"
    assert s.source_type.value == "xlsx"


def test_word_en_excel_labels_nl(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    assert dialog.radio_word.text() == "Word-document"
    assert dialog.radio_excel.text() == "Excel-werkmap"
    assert dialog.word_pick_btn.text() == "Kies Word…"
    assert dialog.excel_pick_btn.text() == "Kies Excel…"


def test_word_en_excel_labels_en(tmp_path):
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="en_US", import_service=service)

    assert dialog.radio_word.text() == "Word document"
    assert dialog.radio_excel.text() == "Excel workbook"
    assert dialog.word_pick_btn.text() == "Choose Word…"
    assert dialog.excel_pick_btn.text() == "Choose Excel…"


def test_taalwissel_behoudt_word_selectie(tmp_path):
    """Bij een taalwissel blijft de Word-selectie staan."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)
    dialog.radio_word.setChecked(True)

    dialog.taal = "en_US"
    dialog._apply_language()

    assert dialog.radio_word.isChecked() is True
    assert dialog.radio_word.text() == "Word document"


# ============================================================================
# v1.7.0 — succesmelding na import (fase 6C)
# ============================================================================

def test_importeren_toont_succesmelding(tmp_path, monkeypatch):
    """Na een geslaagde import roept de wizard _toon_succesmelding aan.

    We patchen de Python-methode op de dialoog-instantie, niet
    QMessageBox.information. Dat laatste is een C++-staticmethod die
    zich niet betrouwbaar laat monkeypatchen en de test zou laten
    hangen op een modale dialoog.
    """
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    ontvangen = {"count": 0, "titel": None}

    def _fake_succesmelding(titel):
        ontvangen["count"] += 1
        ontvangen["titel"] = titel

    monkeypatch.setattr(dialog, "_toon_succesmelding", _fake_succesmelding)

    dialog._perform_import()

    assert ontvangen["count"] == 1
    # De PDF-titel is afgeleid van de bestandsnaam (doc).
    assert ontvangen["titel"] == "doc"
    assert len(service.list_sources()) == 1


def test_keep_actie_toont_geen_succesmelding(tmp_path, monkeypatch):
    """Bij KEEP (changed=False) roept de wizard _toon_succesmelding niet aan."""
    _app()
    service = ImportService(catalog_path=tmp_path / "cat.json")
    dialog = ImportWizardDialog(taal="nl_NL", import_service=service)

    pdf_pad = _maak_eenvoudige_pdf(tmp_path / "doc.pdf", titel="Mijn PDF")
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        staticmethod(lambda *a, **kw: (str(pdf_pad), "PDF (*.pdf)")),
    )
    monkeypatch.setattr(wizard_module, "sla_instellingen_op", lambda _: None)
    dialog._pick_pdf()

    # Eerste import om een bestaande bron te creëren
    from app.documentation.pdf_import import import_pdf
    import_pdf(
        pdf_pad,
        import_service=service,
        sources_dir=tmp_path / "sources",
    )

    # Popup: KEEP
    from app.gui.dialogs import duplicate_source_dialog as dup_module
    monkeypatch.setattr(
        dup_module.DuplicateSourceDialog,
        "vraag_actie",
        staticmethod(lambda **kw: DuplicateAction.KEEP),
    )

    ontvangen = {"count": 0}

    def _fake_succesmelding(titel):
        ontvangen["count"] += 1

    monkeypatch.setattr(dialog, "_toon_succesmelding", _fake_succesmelding)

    dialog._perform_import()

    assert ontvangen["count"] == 0
    # De bestaande bron blijft ongewijzigd.
    assert len(service.list_sources()) == 1