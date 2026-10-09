"""
================================================================================
Module:     app/documentation/xlsx_extract.py
Project:    Electronics Diagnostic Tool Hub / ESR Tester (Windows)
Versie:     1.0.1
Datum:      2026-10-09
Auteur:     Bart Bossuyt

Doel:       GUI-onafhankelijke Excel-hulpfuncties voor de import-wizard:
            metadata lezen, celinhoud extraheren, SHA-256 berekenen en het
            bronbestand kopiëren naar de gebruikersmap. Geen Qt, geen
            netwerk, geen assessment.

            Spiegel van pdf_extract.py en docx_extract.py. Deelt bewust
            dezelfde API-vorm.

Wijzigingen:
  v1.0.0 (2026-10-09)  Eerste versie (fase 6A): extract_xlsx_metadata,
                       extract_xlsx_text, compute_file_hash,
                       copy_xlsx_to_sources, default_sources_dir,
                       rename_existing_xlsx_to_old, restore_old_xlsx.
  v1.0.1 (2026-10-09)  extract_xlsx_text: elk werkblad krijgt een header,
                       ook als het blad leeg is. Maakt de structuur van
                       de werkmap zichtbaar voor latere AI-extractie.
================================================================================
"""

from __future__ import annotations

import hashlib
import os
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class XlsxExtractError(ValueError):
    """Fout bij het lezen, hashen of kopiëren van een Excel-bron."""


@dataclass(frozen=True, slots=True)
class XlsxMetadata:
    """Read-only metadata van een Excel-werkmap.

    Alle velden zijn optioneel behalve sheet_count en sheet_names.
    """

    sheet_count: int
    sheet_names: tuple[str, ...]
    title: Optional[str] = None
    author: Optional[str] = None
    subject: Optional[str] = None
    created: Optional[str] = None
    modified: Optional[str] = None


def default_sources_dir() -> Path:
    """Standaardmap voor gekopieerde bronbestanden.

    Identiek aan pdf_extract.default_sources_dir().
    """
    base = os.environ.get("LOCALAPPDATA")
    if not base:
        base = str(Path.home() / "AppData" / "Local")
    return (
        Path(base)
        / "ElectronicsDiagnosticToolHub"
        / "documentation"
        / "sources"
    )


def extract_xlsx_metadata(pad: Path) -> XlsxMetadata:
    """Lees metadata uit een Excel-werkmap.

    Faalt met XlsxExtractError bij een ontbrekend, onleesbaar of
    corrupt bestand.
    """
    pad = Path(pad)
    if not pad.exists():
        raise XlsxExtractError(f"Excel-bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise XlsxExtractError(f"pad is geen bestand: {pad}")

    try:
        # read_only=True: alleen lezen, geen wijzigingen. data_only=False
        # zodat formules leesbaar blijven als tekst.
        wb = load_workbook(str(pad), read_only=True, data_only=False)
    except InvalidFileException as exc:
        raise XlsxExtractError(
            f"Excel-bestand onleesbaar (geen geldig .xlsx-pakket): {exc}"
        ) from exc
    except OSError as exc:
        raise XlsxExtractError(
            f"Excel-bestand kan niet geopend worden: {exc}"
        ) from exc
    except Exception as exc:  # noqa: BLE001 — openpyxl gooit diverse types
        raise XlsxExtractError(
            f"Excel-bestand kon niet gelezen worden: {exc}"
        ) from exc

    try:
        eigenschappen = wb.properties
        bladen = tuple(wb.sheetnames)

        def _tekst(waarde) -> Optional[str]:
            if waarde is None:
                return None
            tekst = str(waarde).strip()
            return tekst or None

        def _datum(waarde) -> Optional[str]:
            if waarde is None:
                return None
            try:
                return waarde.isoformat()
            except AttributeError:
                tekst = str(waarde).strip()
                return tekst or None

        return XlsxMetadata(
            sheet_count=len(bladen),
            sheet_names=bladen,
            title=_tekst(getattr(eigenschappen, "title", None)),
            author=_tekst(getattr(eigenschappen, "creator", None)),
            subject=_tekst(getattr(eigenschappen, "subject", None)),
            created=_datum(getattr(eigenschappen, "created", None)),
            modified=_datum(getattr(eigenschappen, "modified", None)),
        )
    finally:
        wb.close()


def extract_xlsx_text(pad: Path) -> str:
    """Extraheer alle celwaarden als tekst, per werkblad gescheiden.

    Formules worden als formule-tekst weergegeven (bv. "=SUM(A1:A5)"),
    niet als berekend resultaat. Lege cellen worden overgeslagen.

    Elk werkblad krijgt een header ``# <bladnaam>``, ook als het blad
    leeg is. Zo blijft de structuur van de werkmap zichtbaar voor
    latere AI-extractie (fase 12).

    Voor fase 6 is dit optioneel; de wizard bewaart alleen het bestand.
    """
    pad = Path(pad)
    if not pad.exists():
        raise XlsxExtractError(f"Excel-bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise XlsxExtractError(f"pad is geen bestand: {pad}")

    try:
        wb = load_workbook(str(pad), read_only=True, data_only=False)
    except InvalidFileException as exc:
        raise XlsxExtractError(
            f"Excel-bestand onleesbaar: {exc}"
        ) from exc
    except OSError as exc:
        raise XlsxExtractError(
            f"Excel-bestand kan niet geopend worden: {exc}"
        ) from exc
    except Exception as exc:  # noqa: BLE001
        raise XlsxExtractError(
            f"Excel-bestand kon niet gelezen worden: {exc}"
        ) from exc

    try:
        blokken: list[str] = []
        for naam in wb.sheetnames:
            ws = wb[naam]
            regels: list[str] = []
            for rij in ws.iter_rows(values_only=False):
                waarden = [
                    str(cel.value) for cel in rij if cel.value is not None
                ]
                if waarden:
                    regels.append(" | ".join(waarden))
            # Header altijd, ook bij een leeg blad.
            if regels:
                blokken.append(f"# {naam}\n" + "\n".join(regels))
            else:
                blokken.append(f"# {naam}")
        return "\n\n".join(blokken)
    finally:
        wb.close()


def compute_file_hash(pad: Path, *, blokgrootte: int = 65536) -> str:
    """Bereken de SHA-256 van een bestand, blok per blok."""
    pad = Path(pad)
    if not pad.exists():
        raise XlsxExtractError(f"bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise XlsxExtractError(f"pad is geen bestand: {pad}")

    hasher = hashlib.sha256()
    try:
        with pad.open("rb") as handle:
            while True:
                blok = handle.read(blokgrootte)
                if not blok:
                    break
                hasher.update(blok)
    except OSError as exc:
        raise XlsxExtractError(f"hash kon niet berekend worden: {exc}") from exc

    return hasher.hexdigest()


def copy_xlsx_to_sources(
    bron_pad: Path,
    source_id: str,
    *,
    sources_dir: Optional[Path] = None,
) -> Path:
    """Kopieer een Excel-bestand naar de gebruikersmap en geef het doelpad terug.

    Doelnaam is <source_id>.xlsx. Weigert een bestaand doel te overschrijven.
    """
    bron_pad = Path(bron_pad)
    if not bron_pad.exists():
        raise XlsxExtractError(f"bronbestand niet gevonden: {bron_pad}")
    if not bron_pad.is_file():
        raise XlsxExtractError(f"bronpad is geen bestand: {bron_pad}")
    if not source_id:
        raise XlsxExtractError("source_id mag niet leeg zijn")

    doel_dir = Path(sources_dir) if sources_dir else default_sources_dir()

    try:
        doel_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise XlsxExtractError(
            f"doelmap kon niet aangemaakt worden: {doel_dir} ({exc})"
        ) from exc

    doel_pad = doel_dir / f"{source_id}.xlsx"

    if doel_pad.exists():
        raise XlsxExtractError(
            f"doelbestand bestaat al, weigert te overschrijven: {doel_pad}"
        )

    try:
        shutil.copy2(bron_pad, doel_pad)
    except OSError as exc:
        raise XlsxExtractError(f"kopiëren faalde: {exc}") from exc

    return doel_pad


def rename_existing_xlsx_to_old(
    pad: Path,
    *,
    timestamp: str,
) -> Optional[Path]:
    """Hernoem een bestaand Excel-bronbestand naar .oud-<timestamp>."""
    pad = Path(pad)
    if not pad.exists():
        return None
    if not pad.is_file():
        raise XlsxExtractError(f"pad is geen bestand: {pad}")

    oud_pad = pad.with_name(pad.name + f".oud-{timestamp}")
    if oud_pad.exists():
        raise XlsxExtractError(
            f"oud-bestand bestaat al, weiger te overschrijven: {oud_pad}"
        )

    try:
        os.rename(pad, oud_pad)
    except OSError as exc:
        raise XlsxExtractError(
            f"hernoemen naar oud-bestand faalde: {exc}"
        ) from exc

    return oud_pad


def restore_old_xlsx(old_pad: Path, origineel_pad: Path) -> None:
    """Zet een .oud-<timestamp> bestand terug op zijn originele plaats."""
    old_pad = Path(old_pad)
    origineel_pad = Path(origineel_pad)

    if not old_pad.exists():
        raise XlsxExtractError(
            f"oud-bestand niet gevonden voor rollback: {old_pad}"
        )
    if origineel_pad.exists():
        raise XlsxExtractError(
            f"rollback-doel bestaat al, weiger te overschrijven: {origineel_pad}"
        )

    try:
        os.rename(old_pad, origineel_pad)
    except OSError as exc:
        raise XlsxExtractError(
            f"rollback van oud-bestand faalde: {exc}"
        ) from exc