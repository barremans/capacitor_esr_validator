"""
================================================================================
Module:     app/documentation/docx_extract.py
Project:    Electronics Diagnostic Tool Hub / ESR Tester (Windows)
Versie:     1.0.0
Datum:      2026-10-09
Auteur:     Bart Bossuyt

Doel:       GUI-onafhankelijke Word-hulpfuncties voor de import-wizard:
            metadata lezen, tekst extraheren, SHA-256 berekenen en het
            bronbestand kopiëren naar de gebruikersmap. Geen Qt, geen
            netwerk, geen assessment.

            Spiegel van pdf_extract.py. Deelt bewust dezelfde API-vorm
            (default_sources_dir, compute_file_hash, rename-to-old,
            restore_old) zodat de wizard en de orkestratielaag op
            dezelfde manier met Word-bestanden kunnen omgaan.

Wijzigingen:
  v1.0.0 (2026-10-09)  Eerste versie (fase 6A): extract_docx_metadata,
                       extract_docx_text, compute_file_hash,
                       copy_docx_to_sources, default_sources_dir,
                       rename_existing_docx_to_old, restore_old_docx.
================================================================================
"""

from __future__ import annotations

import hashlib
import os
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from docx import Document
from docx.opc.exceptions import PackageNotFoundError


class DocxExtractError(ValueError):
    """Fout bij het lezen, hashen of kopiëren van een Word-bron."""


@dataclass(frozen=True, slots=True)
class DocxMetadata:
    """Read-only metadata van een Word-bestand.

    Alle velden zijn optioneel behalve paragraph_count: Word-bestanden
    zijn niet verplicht om titel, auteur of aanmaakdatum te bevatten.
    """

    paragraph_count: int
    title: Optional[str] = None
    author: Optional[str] = None
    subject: Optional[str] = None
    created: Optional[str] = None
    modified: Optional[str] = None
    last_modified_by: Optional[str] = None


def default_sources_dir() -> Path:
    """Standaardmap voor gekopieerde bronbestanden.

    Identiek aan pdf_extract.default_sources_dir(): alle lokale
    bronbestanden (PDF, DOCX, XLSX) leven in dezelfde map onder
    %LOCALAPPDATA%.
    """
    base = os.environ.get("LOCALAPPDATA")
    if not base:
        base = str(Path.home() / "AppData" / "Local")
    return (
        Path(base)
        / "ElectronicsDiagnosticToolHub"
        / "documentation"
        / "sources"
    )


def extract_docx_metadata(pad: Path) -> DocxMetadata:
    """Lees metadata uit een Word-bestand.

    Faalt met DocxExtractError bij een ontbrekend, onleesbaar of
    corrupt bestand.
    """
    pad = Path(pad)
    if not pad.exists():
        raise DocxExtractError(f"Word-bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise DocxExtractError(f"pad is geen bestand: {pad}")

    try:
        document = Document(str(pad))
    except PackageNotFoundError as exc:
        raise DocxExtractError(
            f"Word-bestand onleesbaar (geen geldig .docx-pakket): {exc}"
        ) from exc
    except OSError as exc:
        raise DocxExtractError(
            f"Word-bestand kan niet geopend worden: {exc}"
        ) from exc

    core = document.core_properties

    def _tekst(waarde: Optional[str]) -> Optional[str]:
        if waarde is None:
            return None
        tekst = str(waarde).strip()
        return tekst or None

    def _datum(waarde) -> Optional[str]:
        if waarde is None:
            return None
        try:
            return waarde.isoformat()
        except AttributeError:
            tekst = str(waarde).strip()
            return tekst or None

    return DocxMetadata(
        paragraph_count=len(document.paragraphs),
        title=_tekst(core.title),
        author=_tekst(core.author),
        subject=_tekst(core.subject),
        created=_datum(core.created),
        modified=_datum(core.modified),
        last_modified_by=_tekst(core.last_modified_by),
    )


def extract_docx_text(pad: Path) -> str:
    """Extraheer de ruwe tekst van alle paragrafen, gescheiden door lege regels.

    Geen tabellen, geen opmaak, geen afbeeldingen. Voor fase 6 is dit
    optioneel; de wizard bewaart alleen het bestand. De functie is
    aanwezig zodat een latere fase (12 — AI-extractie) er gebruik van
    kan maken.
    """
    pad = Path(pad)
    if not pad.exists():
        raise DocxExtractError(f"Word-bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise DocxExtractError(f"pad is geen bestand: {pad}")

    try:
        document = Document(str(pad))
    except PackageNotFoundError as exc:
        raise DocxExtractError(
            f"Word-bestand onleesbaar: {exc}"
        ) from exc
    except OSError as exc:
        raise DocxExtractError(
            f"Word-bestand kan niet geopend worden: {exc}"
        ) from exc

    teksten = [p.text.strip() for p in document.paragraphs]
    return "\n\n".join(t for t in teksten if t)


def compute_file_hash(pad: Path, *, blokgrootte: int = 65536) -> str:
    """Bereken de SHA-256 van een bestand, blok per blok.

    Zelfde semantiek als pdf_extract.compute_file_hash.
    """
    pad = Path(pad)
    if not pad.exists():
        raise DocxExtractError(f"bestand niet gevonden: {pad}")
    if not pad.is_file():
        raise DocxExtractError(f"pad is geen bestand: {pad}")

    hasher = hashlib.sha256()
    try:
        with pad.open("rb") as handle:
            while True:
                blok = handle.read(blokgrootte)
                if not blok:
                    break
                hasher.update(blok)
    except OSError as exc:
        raise DocxExtractError(f"hash kon niet berekend worden: {exc}") from exc

    return hasher.hexdigest()


def copy_docx_to_sources(
    bron_pad: Path,
    source_id: str,
    *,
    sources_dir: Optional[Path] = None,
) -> Path:
    """Kopieer een Word-bestand naar de gebruikersmap en geef het doelpad terug.

    Doelnaam is <source_id>.docx. Weigert een bestaand doel te overschrijven.
    """
    bron_pad = Path(bron_pad)
    if not bron_pad.exists():
        raise DocxExtractError(f"bronbestand niet gevonden: {bron_pad}")
    if not bron_pad.is_file():
        raise DocxExtractError(f"bronpad is geen bestand: {bron_pad}")
    if not source_id:
        raise DocxExtractError("source_id mag niet leeg zijn")

    doel_dir = Path(sources_dir) if sources_dir else default_sources_dir()

    try:
        doel_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise DocxExtractError(
            f"doelmap kon niet aangemaakt worden: {doel_dir} ({exc})"
        ) from exc

    doel_pad = doel_dir / f"{source_id}.docx"

    if doel_pad.exists():
        raise DocxExtractError(
            f"doelbestand bestaat al, weigert te overschrijven: {doel_pad}"
        )

    try:
        shutil.copy2(bron_pad, doel_pad)
    except OSError as exc:
        raise DocxExtractError(f"kopiëren faalde: {exc}") from exc

    return doel_pad


def rename_existing_docx_to_old(
    pad: Path,
    *,
    timestamp: str,
) -> Optional[Path]:
    """Hernoem een bestaand Word-bronbestand naar .oud-<timestamp>.

    Wordt gebruikt door 'Overschrijven' om de originele bron te bewaren
    vóór het nieuwe bestand wordt gekopieerd.
    """
    pad = Path(pad)
    if not pad.exists():
        return None
    if not pad.is_file():
        raise DocxExtractError(f"pad is geen bestand: {pad}")

    oud_pad = pad.with_name(pad.name + f".oud-{timestamp}")
    if oud_pad.exists():
        raise DocxExtractError(
            f"oud-bestand bestaat al, weiger te overschrijven: {oud_pad}"
        )

    try:
        os.rename(pad, oud_pad)
    except OSError as exc:
        raise DocxExtractError(
            f"hernoemen naar oud-bestand faalde: {exc}"
        ) from exc

    return oud_pad


def restore_old_docx(old_pad: Path, origineel_pad: Path) -> None:
    """Zet een .oud-<timestamp> bestand terug op zijn originele plaats."""
    old_pad = Path(old_pad)
    origineel_pad = Path(origineel_pad)

    if not old_pad.exists():
        raise DocxExtractError(
            f"oud-bestand niet gevonden voor rollback: {old_pad}"
        )
    if origineel_pad.exists():
        raise DocxExtractError(
            f"rollback-doel bestaat al, weiger te overschrijven: {origineel_pad}"
        )

    try:
        os.rename(old_pad, origineel_pad)
    except OSError as exc:
        raise DocxExtractError(
            f"rollback van oud-bestand faalde: {exc}"
        ) from exc