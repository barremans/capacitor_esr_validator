"""
================================================================================
Module:     app/helpers/history_import.py
Project:    Electronics Diagnostic Tool Hub / ESR Tester (Windows)
Versie:     1.1.0
Datum:      2026-10-27
Auteur:     Bart Bossuyt

Doel:       Pure, GUI- en database-onafhankelijke parser en validator voor
            de historiek-import (C1).

            De parser leest een rij-iterable (kolomkoppen + datarijen) en
            zet die om naar een lijst van ImportKeten-objecten. Elke keten
            bevat de canonieke velden uit DETAIL_EXPORT_HEADERS, gegroepeerd
            per entiteit, plus een natural_key (SHA-256) voor dedup.

            De parser doet GEEN database-IO en GEEN assessment-herberekening.
            Hij is formaat-agnostisch: de CSV- en XLSX-adapters leveren
            rijen, de rest is identiek.

Wijzigingen:
  v1.0.0 (2026-10-27)  Eerste versie. CSV-only via een rij-iterable;
                        natural key volgens de regels van
                        docs/export_import_format.md §6, toegepast op de
                        detail-export-kolommen. tool_key moet
                        ESR_CAPACITOR zijn; andere tool_keys worden
                        overgeslagen met een melding, niet het hele
                        bestand geweigerd.
  v1.1.0 (2026-10-27)  Excel-adapter toegevoegd via openpyxl (read-only).
                        Nieuwe publieke functies lees_csv_rijen() en
                        lees_xlsx_rijen(). parse_bestand() kiest op
                        basis van de extensie. Geen wijziging aan de
                        bestaande parse- of validatielogica.
================================================================================
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence

from app.helpers.history_detail_export import DETAIL_EXPORT_HEADERS


# ---------------------------------------------------------------------------
# Constanten
# ---------------------------------------------------------------------------

TOEGESTANE_TOOL_KEYS = ("ESR_CAPACITOR",)
TOEGESTANE_MEETMETHODEN = ("EX_SITU", "ONE_LEG", "IN_CIRCUIT")
TOEGESTANE_SAMPLE_STATES = (
    "NEW",
    "NOS_UNKNOWN_AGE",
    "USED_WORKING",
    "USED_UNKNOWN",
    "SUSPECT",
    "FAILED_CONFIRMED",
    "REFERENCE_SAMPLE",
    "UNKNOWN",
)
TOEGESTANE_MEETSTATUSSEN = ("NORMAL", "OL", "UNDER_RANGE", "OVER_RANGE", "UNKNOWN")

ONDERSTEUNDE_EXTENSIES = (".csv", ".xlsx")

# Kolommen die als float/int gelezen moeten worden. Leeg blijft None.
NUMERIEKE_KOLOMMEN = frozenset(
    {
        "measured_at_ms",
        "measurement_session_id",
        "component_sample_id",
        "component_model_id",
        "frequency_hz",
        "test_voltage_vrms",
        "temperature_c",
        "residual_voltage_before_v",
        "residual_voltage_after_v",
        "assessment_count",
        "latest_assessment_id",
        "assessed_at_ms",
        "reference_count",
        "reference_value",
        "reference_frequency_hz",
        "reference_temperature_c",
        "reference_test_voltage_vrms",
        "reference_dc_bias_v",
        "esr_nominal_capacitance_f",
        "esr_nominal_capacitance_value",
        "esr_rated_voltage_v",
        "esr_tolerance_lower_pct",
        "esr_tolerance_upper_pct",
        "esr_measured_capacitance_f",
        "esr_measured_esr_ohm",
        "esr_dissipation_factor_d",
        "esr_quality_factor_q",
        "esr_impedance_z_ohm",
        "esr_reactance_x_ohm",
        "esr_capacitance_deviation_pct",
        "esr_esr_factor",
        "esr_derived_esr_from_d_ohm",
    }
)

BOOLEAN_KOLOMMEN = frozenset(
    {
        "power_off_confirmed",
        "discharged_confirmed",
        "out_of_range",
        "open_suspected",
        "short_suspected",
        "unstable_reading",
    }
)

JSON_KOLOMMEN = frozenset(
    {"reasons_json", "warnings_json", "advice_json", "references_json"}
)


# ---------------------------------------------------------------------------
# Dataclasses
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class ImportFout:
    """Eén fout of skip in één rij, met de rij-index (1-based, kop = 1)."""

    rij_index: int
    kolom: str
    bericht: str


@dataclass(frozen=True, slots=True)
class ImportKeten:
    """Eén gevalideerde meetketen, klaar voor de import-service.

    De velden zijn gegroepeerd per entiteit, zodat de service ze direct
    aan save_new_measurement_chain() kan doorgeven. Geen id's, geen
    created_at_ms, geen source_machine_id/source_measurement_id.
    """

    natural_key: str
    tool_key: str
    component_model: dict[str, Any]
    component_sample: dict[str, Any]
    measurement_session: dict[str, Any]
    measurement: dict[str, Any]
    assessment_snapshot: dict[str, Any] | None
    reference_snapshots: tuple[dict[str, Any], ...]


@dataclass(frozen=True, slots=True)
class ImportResultaat:
    """Uitkomst van het parsen van één bestand."""

    ketens: tuple[ImportKeten, ...]
    fouten: tuple[ImportFout, ...]
    totaal_rijen: int
    ontbrekende_kolommen: tuple[str, ...] = ()


class ImportValidatieFout(ValueError):
    """Het bestand als geheel is ongeldig (bv. ontbrekende kolommen)."""


# ---------------------------------------------------------------------------
# Hulpfuncties
# ---------------------------------------------------------------------------


def _normaliseer_kop(kop: Any) -> str:
    """Koppen zijn exact DETAIL_EXPORT_HEADERS; alleen witruimte strippen."""
    if kop is None:
        return ""
    return str(kop).strip()


def _parse_float(waarde: Any) -> float | None:
    if waarde is None:
        return None
    tekst = str(waarde).strip()
    if tekst == "":
        return None
    getal = float(tekst)
    if not math.isfinite(getal):
        raise ValueError("niet-eindige waarde")
    return getal


def _parse_int(waarde: Any) -> int | None:
    if waarde is None:
        return None
    tekst = str(waarde).strip()
    if tekst == "":
        return None
    getal = float(tekst)
    if not math.isfinite(getal):
        raise ValueError("niet-eindige waarde")
    if getal != int(getal):
        raise ValueError("geen geheel getal")
    return int(getal)


def _parse_bool(waarde: Any) -> bool | None:
    if waarde is None:
        return None
    tekst = str(waarde).strip().lower()
    if tekst == "":
        return None
    if tekst in ("true", "1", "ja", "yes"):
        return True
    if tekst in ("false", "0", "nee", "no"):
        return False
    raise ValueError("geen geldige boolean")


def _parse_json(waarde: Any) -> Any:
    if waarde is None:
        return None
    tekst = str(waarde).strip()
    if tekst == "":
        return None
    return json.loads(tekst)


def _tekst(waarde: Any) -> str | None:
    if waarde is None:
        return None
    tekst = str(waarde).strip()
    return tekst or None


# ---------------------------------------------------------------------------
# Canonieke keten + natural key
# ---------------------------------------------------------------------------


def _canonieke_dict(keten: ImportKeten) -> dict[str, Any]:
    """Bouwt de canonieke dict voor de natural key.

    Regels (docs/export_import_format.md §6): alleen de velden uit de
    keten, None weglaten, JSON met sort_keys=True en compacte separators,
    UTF-8, SHA-256. Geen natural_key, geen source_*, geen id's, geen
    created_at_ms.
    """

    def schoon(data: Mapping[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in data.items() if v is not None}

    payload = {
        "component_model": schoon(keten.component_model),
        "component_sample": schoon(keten.component_sample),
        "measurement_session": schoon(keten.measurement_session),
        "measurement": schoon(keten.measurement),
        "assessment_snapshot": (
            schoon(keten.assessment_snapshot)
            if keten.assessment_snapshot is not None
            else None
        ),
        "reference_snapshots": [schoon(ref) for ref in keten.reference_snapshots],
    }
    return payload


def bereken_natural_key(keten: ImportKeten) -> str:
    """SHA-256 hex over de canonieke keten (zonder toestel-specifieke velden)."""
    payload = _canonieke_dict(keten)
    tekst = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(tekst.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Rij → keten
# ---------------------------------------------------------------------------


def _rij_naar_dict(
    kop: Sequence[str],
    waarden: Sequence[Any],
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for index, kolom in enumerate(kop):
        if not kolom:
            continue
        result[kolom] = waarden[index] if index < len(waarden) else None
    return result


def _bouw_keten(
    rij: Mapping[str, Any],
    rij_index: int,
) -> tuple[ImportKeten | None, list[ImportFout]]:
    fouten: list[ImportFout] = []

    def fout(kolom: str, bericht: str) -> None:
        fouten.append(ImportFout(rij_index=rij_index, kolom=kolom, bericht=bericht))

    tool_key = _tekst(rij.get("tool_key"))
    if not tool_key:
        fout("tool_key", "verplicht")
        return None, fouten
    if tool_key not in TOEGESTANE_TOOL_KEYS:
        fout("tool_key", f"niet ondersteund: {tool_key}")
        return None, fouten

    measurement_method = _tekst(rij.get("measurement_method"))
    if not measurement_method or measurement_method not in TOEGESTANE_MEETMETHODEN:
        fout("measurement_method", f"ongeldig: {measurement_method!r}")

    sample_state = _tekst(rij.get("sample_state"))
    if not sample_state or sample_state not in TOEGESTANE_SAMPLE_STATES:
        fout("sample_state", f"ongeldig: {sample_state!r}")

    try:
        measured_at_ms = _parse_int(rij.get("measured_at_ms"))
    except ValueError:
        measured_at_ms = None
    if measured_at_ms is None or measured_at_ms < 0:
        fout("measured_at_ms", "moet een niet-negatieve integer zijn")

    numeriek: dict[str, float | int | None] = {}
    for kolom in NUMERIEKE_KOLOMMEN:
        if kolom == "measured_at_ms":
            continue
        if kolom not in rij:
            continue
        try:
            if kolom in (
                "measurement_session_id",
                "component_sample_id",
                "component_model_id",
                "latest_assessment_id",
                "assessment_count",
                "reference_count",
                "assessed_at_ms",
            ):
                numeriek[kolom] = _parse_int(rij.get(kolom))
            else:
                numeriek[kolom] = _parse_float(rij.get(kolom))
        except ValueError as exc:
            fout(kolom, str(exc))

    booleans: dict[str, bool | None] = {}
    for kolom in BOOLEAN_KOLOMMEN:
        if kolom not in rij:
            continue
        try:
            booleans[kolom] = _parse_bool(rij.get(kolom))
        except ValueError as exc:
            fout(kolom, str(exc))

    json_velden: dict[str, Any] = {}
    for kolom in JSON_KOLOMMEN:
        if kolom not in rij:
            continue
        try:
            json_velden[kolom] = _parse_json(rij.get(kolom))
        except (TypeError, ValueError) as exc:
            fout(kolom, f"ongeldige JSON: {exc}")

    if fouten:
        return None, fouten

    component_model = {
        "manufacturer": _tekst(rij.get("manufacturer")),
        "series": _tekst(rij.get("series")),
        "part_number": _tekst(rij.get("part_number")),
        "technology": _tekst(rij.get("technology")),
        "nominal_capacitance_f": numeriek.get("esr_nominal_capacitance_f"),
        "nominal_capacitance_value": numeriek.get("esr_nominal_capacitance_value"),
        "nominal_capacitance_unit": _tekst(rij.get("esr_nominal_capacitance_unit")),
        "rated_voltage_v": numeriek.get("esr_rated_voltage_v"),
        "voltage_type": _tekst(rij.get("esr_voltage_type")),
        "tolerance_lower_pct": numeriek.get("esr_tolerance_lower_pct"),
        "tolerance_upper_pct": numeriek.get("esr_tolerance_upper_pct"),
    }

    component_sample = {
        "sample_code": _tekst(rij.get("sample_code")),
        "sample_state": sample_state,
    }

    measurement_session = {
        "started_at_ms": measured_at_ms,
        "ended_at_ms": None,
        "operator_name": None,
        "customer": _tekst(rij.get("customer")),
        "project": _tekst(rij.get("project")),
        "installation": _tekst(rij.get("installation")),
        "module_board": _tekst(rij.get("module_board")),
        "component_reference": _tekst(rij.get("component_reference")),
        "notes": None,
    }

    measurement = {
        "measured_at_ms": measured_at_ms,
        "measurement_method": measurement_method,
        "tool_key": tool_key,
        "instrument_key": _tekst(rij.get("instrument_key")),
        "instrument_name": _tekst(rij.get("instrument_name")),
        "instrument_profile_version": _tekst(
            rij.get("instrument_profile_version")
        ),
        "frequency_hz": numeriek.get("frequency_hz"),
        "test_voltage_vrms": numeriek.get("test_voltage_vrms"),
        "temperature_c": numeriek.get("temperature_c"),
        "power_off_confirmed": booleans.get("power_off_confirmed"),
        "discharged_confirmed": booleans.get("discharged_confirmed"),
        "residual_voltage_before_v": numeriek.get("residual_voltage_before_v"),
        "residual_voltage_after_v": numeriek.get("residual_voltage_after_v"),
        "capacitance_f": numeriek.get("esr_measured_capacitance_f"),
        "esr_ohm": numeriek.get("esr_measured_esr_ohm"),
        "dissipation_factor_d": numeriek.get("esr_dissipation_factor_d"),
        "quality_factor_q": numeriek.get("esr_quality_factor_q"),
        "impedance_z_ohm": numeriek.get("esr_impedance_z_ohm"),
        "reactance_x_ohm": numeriek.get("esr_reactance_x_ohm"),
        "out_of_range": booleans.get("out_of_range", False) or False,
        "open_suspected": booleans.get("open_suspected", False) or False,
        "short_suspected": booleans.get("short_suspected", False) or False,
        "unstable_reading": booleans.get("unstable_reading", False) or False,
        "parallel_components_notes": _tekst(rij.get("parallel_components_notes")),
        "mechanical_condition_notes": _tekst(
            rij.get("mechanical_condition_notes")
        ),
        "notes": _tekst(rij.get("measurement_notes")),
        "record_status": _tekst(rij.get("record_status")),
        "invalid_reason": _tekst(rij.get("invalid_reason")),
    }

    assessment_snapshot = None
    assessed_at_ms = numeriek.get("assessed_at_ms")
    if assessed_at_ms is not None:
        assessment_snapshot = {
            "assessed_at_ms": int(assessed_at_ms),
            "engine_version": _tekst(rij.get("assessment_engine_version")),
            "capacitance_status": _tekst(rij.get("esr_capacitance_status")),
            "capacitance_deviation_pct": numeriek.get(
                "esr_capacitance_deviation_pct"
            ),
            "esr_status": _tekst(rij.get("esr_esr_status")),
            "esr_factor": numeriek.get("esr_esr_factor"),
            "consistency_status": _tekst(rij.get("esr_consistency_status")),
            "derived_esr_from_d_ohm": numeriek.get("esr_derived_esr_from_d_ohm"),
            "reliability_level": _tekst(rij.get("reliability_level")),
            "final_status": _tekst(rij.get("final_status")),
            "reasons_json": (
                json.dumps(json_velden["reasons_json"], ensure_ascii=False)
                if json_velden.get("reasons_json") is not None
                else None
            ),
            "warnings_json": (
                json.dumps(json_velden["warnings_json"], ensure_ascii=False)
                if json_velden.get("warnings_json") is not None
                else None
            ),
            "advice_json": (
                json.dumps(json_velden["advice_json"], ensure_ascii=False)
                if json_velden.get("advice_json") is not None
                else None
            ),
        }

    reference_snapshots: list[dict[str, Any]] = []
    refs_raw = json_velden.get("references_json")
    if isinstance(refs_raw, list):
        for ref in refs_raw:
            if not isinstance(ref, dict):
                continue
            reference_snapshots.append(
                {
                    "reference_role": ref.get("reference_role"),
                    "reference_level": ref.get("reference_level"),
                    "reference_type": ref.get("reference_type"),
                    "manufacturer": ref.get("manufacturer"),
                    "series": ref.get("series"),
                    "part_number": ref.get("part_number"),
                    "reference_value": ref.get("reference_value"),
                    "reference_unit": ref.get("reference_unit"),
                    "frequency_hz": ref.get("frequency_hz"),
                    "temperature_c": ref.get("temperature_c"),
                    "test_voltage_vrms": ref.get("test_voltage_vrms"),
                    "dc_bias_v": ref.get("dc_bias_v"),
                    "value_kind": ref.get("value_kind"),
                    "source_name": ref.get("source_name"),
                    "source_document": ref.get("source_document"),
                    "source_version": ref.get("source_version"),
                    "source_data_version": ref.get("source_data_version"),
                    "source_entry_id": ref.get("source_entry_id"),
                    "source_hash": ref.get("source_hash"),
                    "snapshot_json": ref.get("snapshot_json"),
                }
            )

    keten_zonder_sleutel = ImportKeten(
        natural_key="",
        tool_key=tool_key,
        component_model=component_model,
        component_sample=component_sample,
        measurement_session=measurement_session,
        measurement=measurement,
        assessment_snapshot=assessment_snapshot,
        reference_snapshots=tuple(reference_snapshots),
    )
    natural_key = bereken_natural_key(keten_zonder_sleutel)

    return (
        ImportKeten(
            natural_key=natural_key,
            tool_key=tool_key,
            component_model=component_model,
            component_sample=component_sample,
            measurement_session=measurement_session,
            measurement=measurement,
            assessment_snapshot=assessment_snapshot,
            reference_snapshots=tuple(reference_snapshots),
        ),
        fouten,
    )


# ---------------------------------------------------------------------------
# Bestandsadapters
# ---------------------------------------------------------------------------


def lees_csv_rijen(pad: str | Path) -> Iterator[list[str]]:
    """Leest een UTF-8 (BOM of geen BOM) CSV met puntkomma-scheiding."""
    path = Path(pad)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=";")
        for rij in reader:
            yield rij


def lees_xlsx_rijen(pad: str | Path) -> Iterator[list[Any]]:
    """Leest de eerste sheet van een .xlsx als rijen; openpyxl read-only."""
    from openpyxl import load_workbook

    path = Path(pad)
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return
        for rij in sheet.iter_rows(values_only=True):
            yield list(rij)
    finally:
        workbook.close()


def lees_bestand_rijen(pad: str | Path) -> Iterator[list[Any]]:
    """Kiest de adapter op basis van de extensie (.csv of .xlsx)."""
    path = Path(pad)
    extensie = path.suffix.lower()
    if extensie == ".csv":
        return lees_csv_rijen(path)
    if extensie == ".xlsx":
        return lees_xlsx_rijen(path)
    raise ImportValidatieFout(
        f"Niet-ondersteund bestandstype: {extensie or '(geen)'}. "
        f"Gebruik een van: {', '.join(ONDERSTEUNDE_EXTENSIES)}."
    )


# ---------------------------------------------------------------------------
# Publieke API
# ---------------------------------------------------------------------------


def parse_rijen(rijen: Iterable[Sequence[Any]]) -> ImportResultaat:
    """Parseert een rij-iterable (kop + datarijen) naar ImportResultaat.

    De eerste rij is de kop. Ontbrekende verplichte kolommen → het hele
    bestand wordt geweigerd (ImportValidatieFout). Per datarij worden
    fouten verzameld; een rij met fouten levert geen keten op, maar
    blokkeert de andere rijen niet.
    """
    iterator = iter(rijen)
    try:
        kop_raw = next(iterator)
    except StopIteration:
        raise ImportValidatieFout("Het bestand is leeg.") from None

    kop = [_normaliseer_kop(item) for item in kop_raw]
    ontbrekend = [kolom for kolom in DETAIL_EXPORT_HEADERS if kolom not in kop]
    if ontbrekend:
        raise ImportValidatieFout(
            "Ontbrekende verplichte kolommen: " + ", ".join(ontbrekend)
        )

    ketens: list[ImportKeten] = []
    fouten: list[ImportFout] = []
    gezien: set[str] = set()
    totaal = 0
    for rij_index, waarden in enumerate(iterator, start=2):
        totaal += 1
        rij = _rij_naar_dict(kop, list(waarden))
        keten, rij_fouten = _bouw_keten(rij, rij_index)
        if rij_fouten:
            fouten.extend(rij_fouten)
            continue
        if keten is None:
            continue
        if keten.natural_key in gezien:
            fouten.append(
                ImportFout(
                    rij_index=rij_index,
                    kolom="natural_key",
                    bericht="duplicate binnen hetzelfde bestand",
                )
            )
            continue
        gezien.add(keten.natural_key)
        ketens.append(keten)

    return ImportResultaat(
        ketens=tuple(ketens),
        fouten=tuple(fouten),
        totaal_rijen=totaal,
        ontbrekende_kolommen=tuple(ontbrekend),
    )


def parse_bestand(pad: str | Path) -> ImportResultaat:
    """Parseert een .csv- of .xlsx-bestand naar ImportResultaat."""
    return parse_rijen(lees_bestand_rijen(pad))